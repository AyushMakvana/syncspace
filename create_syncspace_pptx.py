import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Exact Color Palette matching the PlaceIntel reference template
    BG_COLOR = RGBColor(248, 250, 248)       # #F8FAF8 (Clean off-white / pale mint background)
    GREEN_PRIMARY = RGBColor(59, 105, 87)    # #3B6957 (Dark Forest Green badge & line color)
    TEXT_DARK = RGBColor(24, 34, 48)         # #182230 (Dark Slate Navy headings)
    TEXT_MUTED = RGBColor(71, 85, 105)       # #475569 (Slate Gray body text)
    DIVIDER_COLOR = RGBColor(203, 213, 225)  # #CBD5E1 (Soft line color)
    CARD_BG = RGBColor(238, 244, 241)        # #EEF4F1 (Pale Mint Tint card fill)
    CARD_BORDER = RGBColor(210, 224, 216)    # #D2E0D8 (Card border)
    WHITE = RGBColor(255, 255, 255)

    FONT_MAIN = "Segoe UI"

    def set_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, category_text, title_text, slide_num):
        set_bg(slide)

        # Top Badge Pill
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(2.3), Inches(0.38))
        badge.fill.solid()
        badge.fill.fore_color.rgb = GREEN_PRIMARY
        badge.line.color.rgb = GREEN_PRIMARY
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.name = FONT_MAIN
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

        # Slide Title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(10), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_MAIN
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        # Top Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = GREEN_PRIMARY
        line.line.color.rgb = GREEN_PRIMARY

        # Footer
        ftb = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        p = ftb.text_frame.paragraphs[0]
        p.text = f"SYNCSPACE · {category_text.upper()}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = GREEN_PRIMARY

        p2 = ftb.text_frame.add_paragraph()
        p2.text = f"{slide_num:02d} / 14"
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MUTED
        p2.alignment = PP_ALIGN.RIGHT

    # ==================== SLIDE 1: Title Slide ====================
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1)

    badge1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(2.5), Inches(0.45))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = GREEN_PRIMARY
    badge1.line.color.rgb = GREEN_PRIMARY
    tf1 = badge1.text_frame
    p = tf1.paragraphs[0]
    p.text = "PROJECT SHOWCASE"
    p.font.name = FONT_MAIN
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(10), Inches(1.5))
    p = tb1.text_frame.paragraphs[0]
    p.text = "SyncSpace"
    p.font.name = FONT_MAIN
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p_sub = tb1.text_frame.add_paragraph()
    p_sub.text = "Real-Time Virtual Watch Party & Collaborative Platform"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = TEXT_MUTED

    line1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.6), Inches(7.5), Inches(0.02))
    line1.fill.solid()
    line1.fill.fore_color.rgb = DIVIDER_COLOR
    line1.line.color.rgb = DIVIDER_COLOR

    tb_info = slide1.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(10), Inches(2.5))
    tf_info = tb_info.text_frame

    p = tf_info.paragraphs[0]
    p.text = "REVIEW 1"
    p.font.name = FONT_MAIN
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf_info.add_paragraph()
    p.text = "Department of Computer Engineering\nCHARUSAT University"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(16)

    p = tf_info.add_paragraph()
    p.text = "Team Members"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf_info.add_paragraph()
    p.text = "Ayush Makvana (Lead Developer)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_MUTED

    ftb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
    p = ftb1.text_frame.paragraphs[0]
    p.text = "SYNCSPACE · CHARUSAT UNIVERSITY"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY

    # ==================== SLIDE 2: Problem Statement ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem", "Problem Statement", 2)

    problems = [
        ("01", "Distance Barrier", "Friends and study partners live in different locations and cannot watch videos together easily."),
        ("02", "Timing Mismatch", "Counting down '3-2-1 play' manually fails due to internet lag, causing video spoilers."),
        ("03", "Hard Setup Process", "Tools like Teleparty force everyone to install extensions, use desktop Chrome, and make paid accounts."),
        ("04", "Separated Chatting", "Texting on WhatsApp while watching ruins the fun; existing sites lack a real shared virtual sofa.")
    ]

    top_pos = 1.9
    for num, title, desc in problems:
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(6.5), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = f"{num}  {title}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(2)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)  # Increased to 12pt!
        p_desc.font.color.rgb = TEXT_MUTED

        top_pos += 1.2

    # Core Need Box Right
    cbox = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.6), Inches(1.9), Inches(4.9), Inches(4.7))
    cbox.fill.solid()
    cbox.fill.fore_color.rgb = CARD_BG
    cbox.line.color.rgb = CARD_BORDER
    tf = cbox.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.4)
    tf.margin_right = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "CORE NEED"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY
    p.space_after = Pt(16)

    p = tf.add_paragraph()
    p.text = "An easy web app is needed so anyone can join watch parties instantly with zero setup, super fast live chat, and a visual sofa room."
    p.font.name = FONT_MAIN
    p.font.size = Pt(19)  # Large font
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    # ==================== SLIDE 3: Existing System Analysis ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Diagnosis", "Existing System Analysis", 3)

    rows = 6
    cols = 5
    table_shape = slide3.shapes.add_table(rows, cols, Inches(0.8), Inches(1.9), Inches(11.733), Inches(4.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(2.3)
    table.columns[2].width = Inches(2.3)
    table.columns[3].width = Inches(2.3)
    table.columns[4].width = Inches(2.333)

    headers = ["Feature / Parameter", "Teleparty (Netflix Party)", "Discord Watch Together", "Kosmi.io", "SyncSpace (Proposed)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN_PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT_MAIN
            p.font.size = Pt(12)  # 12pt header
            p.font.bold = True
            p.font.color.rgb = WHITE

    data = [
        ["Setup Ease", "Hard (Needs Extension)", "Medium (Needs App/Login)", "Easy (Web Browser)", "Super Easy (Instant Web/Code)"],
        ["User Login", "Paid streaming login", "Discord account required", "Guest option available", "Google OAuth / Email Login"],
        ["Member Visuals", "Simple avatar list", "Voice grid box", "2D Virtual Rooms", "Liquid Glass Lounge Sofa"],
        ["Chat Speed", "Slow (500ms – 1200ms)", "Medium (200ms – 500ms)", "Medium (150ms – 300ms)", "Super Fast (< 100ms)"],
        ["Device Support", "Desktop Chrome only", "Discord Desktop/Web", "Modern Browsers", "100% Mobile & Web Browsers"]
    ]

    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_MAIN
                p.font.size = Pt(12)  # Increased to 12pt!
                p.font.color.rgb = GREEN_PRIMARY if c_idx == 4 else TEXT_DARK
                if c_idx == 4:
                    p.font.bold = True

    tb = slide3.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "Summary: Existing apps require extra setup or app downloads. SyncSpace works directly in any browser with instant joining and instant chat."
    p.font.name = FONT_MAIN
    p.font.size = Pt(12)  # Increased to 12pt!
    p.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 4: Objectives & Scope ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Direction", "Project Objectives & Scope", 4)

    card_w = Inches(3.7)
    card_h = Inches(4.7)

    # Col 1: Objectives
    c1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), card_w, card_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.25)

    p = tf.paragraphs[0]
    p.text = "01\nObjectives (Clear Goals)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(12)

    objs = [
        "Fast Chat Goal: Deliver instant live chat & member updates in under 100ms.",
        "Zero Setup Goal: Allow users to join immediately using a simple room code (AYUSHMAKVANA-2779) or link.",
        "High Reliability: Maintain 99.9% login and database uptime using Google Firebase."
    ]
    for o in objs:
        p = tf.add_paragraph()
        p.text = f"• {o}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt font!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # Col 2: Current Scope
    c2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.9), card_w, card_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tf = c2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.25)

    p = tf.paragraphs[0]
    p.text = "02\nCurrent Scope (Completed)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(12)

    curr = [
        "Modern dark liquid-glass website interface.",
        "Firebase Login (Google 1-click popup + Email & Password).",
        "Instant room creation & room code joining.",
        "Real-time lounge sofa showing online friends & live chat box."
    ]
    for c in curr:
        p = tf.add_paragraph()
        p.text = f"• {c}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt font!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # Col 3: Future Scope & Exclusions
    c3 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.9), card_w, card_h)
    c3.fill.solid()
    c3.fill.fore_color.rgb = CARD_BG
    c3.line.color.rgb = CARD_BORDER
    tf = c3.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.25)

    p = tf.paragraphs[0]
    p.text = "03\nNext Scope & Exclusions"
    p.font.name = FONT_MAIN
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(12)

    futs = [
        "Next Phase: YouTube search & synchronized video player (play/pause/scrub).",
        "Next Phase: WebRTC camera and microphone live stream.",
        "Excluded Item 1: Native mobile app downloads (.apk/.ipa) excluded for now.",
        "Excluded Item 2: Paid subscription billing payment features excluded."
    ]
    for f in futs:
        p = tf.add_paragraph()
        p.text = f"• {f}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt font!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # ==================== SLIDE 5: Proposed Solution ====================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Response", "Proposed Solution Architecture", 5)

    sol_w = Inches(2.7)
    sol_h = Inches(2.2)

    solutions = [
        ("01", "Firebase Login", "Google 1-click popup login & Email account creation."),
        ("02", "Room Code Finder", "Typing AYUSHMAKVANA-2779 connects you to the right room."),
        ("03", "Real-Time Sofa", "Displays all active members on sofa with host crown icons."),
        ("04", "Instant Live Chat", "Super fast live messaging working on all browser tabs."),
        ("05", "YouTube Search", "[Next] Search YouTube videos & load playlists directly."),
        ("06", "Camera & Mic", "[Next] Talk and see friends live via WebRTC video stream."),
        ("07", "Video Sync Player", "[Next] Video play, pause & seek syncs for everyone in room."),
        ("08", "Liquid Glass Design", "Modern, smooth glass design with clean animations.")
    ]

    for idx, (num, title, desc) in enumerate(solutions):
        r = idx // 4
        c = idx % 4
        x = Inches(0.8 + c * 2.95)
        y = Inches(1.9 + r * 2.4)

        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, sol_w, sol_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = f"{num}  {title}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)  # Increased to 12pt!
        p_desc.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 6: Target Users ====================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Orientation", "Project Overview & Target Users", 6)

    p_w = Inches(5.6)
    p_h = Inches(2.2)

    personas = [
        ("01", "Rahul Patel (College Student)", "Wants to host weekend movie nights with high school friends across different cities.", "Shares 1-click room code; friends join instantly in browser with zero setup."),
        ("02", "Ananya Sharma (Study Group Lead)", "Coordinates online lecture review sessions for university exam preparation.", "Synchronized lecture watching with live side-chat to ask questions."),
        ("03", "Dev Shah (Media Streamer)", "Hosts watch-party events for tech launch videos with online subscribers.", "Host controls video playback while community members watch together live."),
        ("04", "University Student Clubs", "Technical and cultural student clubs hosting online webinars and movie nights.", "Easy group rooms with visual member avatars and room codes.")
    ]

    for idx, (num, title, scenario, value) in enumerate(personas):
        r = idx // 2
        c = idx % 2
        x = Inches(0.8 + c * 6.0)
        y = Inches(1.9 + r * 2.4)

        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, p_w, p_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = f"{num}  {title}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = f"Goal: {scenario}"
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(12)  # 12pt!
        p1.font.color.rgb = TEXT_MUTED
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = f"Solution: {value}"
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(12)  # 12pt!
        p2.font.bold = True
        p2.font.color.rgb = GREEN_PRIMARY

    # ==================== SLIDE 7: Requirements ====================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Requirements", "Functional & Non-Functional Requirements", 7)

    req_w = Inches(5.6)
    req_h = Inches(4.7)

    # FR Card
    fr_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), req_w, req_h)
    fr_card.fill.solid()
    fr_card.fill.fore_color.rgb = CARD_BG
    fr_card.line.color.rgb = CARD_BORDER
    tf = fr_card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)

    p = tf.paragraphs[0]
    p.text = "SYSTEM FEATURES (FUNCTIONAL)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY
    p.space_after = Pt(10)

    frs = [
        "User Login: Google 1-click popup login and Email account creation.",
        "Room Join: Instant room creation & room code search (AYUSHMAKVANA-2779).",
        "Member Sofa: Live visual sofa showing online users and room host status.",
        "Live Chat: Instant text messaging with user name and time stamps.",
        "Video Sync [Next]: YouTube video player that stays in sync for all users.",
        "Camera/Mic [Next]: WebRTC video and audio stream controls."
    ]
    for item in frs:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # NFR Card
    nfr_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), req_w, req_h)
    nfr_card.fill.solid()
    nfr_card.fill.fore_color.rgb = CARD_BG
    nfr_card.line.color.rgb = CARD_BORDER
    tf = nfr_card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)

    p = tf.paragraphs[0]
    p.text = "QUALITY STANDARDS (NON-FUNCTIONAL)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(10)

    nfrs = [
        "Speed: Super fast messaging & presence updates under 100 milliseconds.",
        "Reliability: 99.9% uptime for logins and rooms using Google Firebase.",
        "Design: Modern liquid-glass dark theme with responsive controls.",
        "Compatibility: Works smoothly across Chrome, Edge, Firefox, & Mobile.",
        "Safety: Safe memory handling to prevent web browser crashing on reload."
    ]
    for item in nfrs:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # ==================== SLIDE 8: Architecture ====================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Structure", "High Level System Architecture", 8)

    diag = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.7))
    diag.fill.solid()
    diag.fill.fore_color.rgb = CARD_BG
    diag.line.color.rgb = CARD_BORDER
    tf = diag.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "SYSTEM LAYERS"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY
    p.space_after = Pt(16)

    layers = [
        "1. Frontend User Layer (Next.js 15 + React 19)\n   Liquid Glass UI · Room Screen · Lounge Sofa · Chat Box",
        "2. Realtime Sync Layer\n   Local Tab Sync (0ms) · Google Firestore Cloud Sync (<100ms)",
        "3. Cloud Services Layer\n   Firebase Auth · Cloud Database · YouTube API v3 · WebRTC"
    ]
    for l in layers:
        p = tf.add_paragraph()
        p.text = l
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    highlights = [
        ("Dual Sync Method", "Combines local tab messaging (0ms) with cloud database sync (<100ms) so chat is instant across all devices."),
        ("Safe Memory Cache", "Uses smart memory cache in Firestore to prevent database lock errors when refreshing browser tabs."),
        ("Room Code Mapper", "Maps user room codes (AYUSHMAKVANA-2779) directly to the right watch party room.")
    ]
    top_pos = 1.9
    for title, desc in highlights:
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(top_pos), Inches(5.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)  # 12pt!
        p_desc.font.color.rgb = TEXT_MUTED

        top_pos += 1.65

    # ==================== SLIDE 9: User Flow ====================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Student Activity Diagram", "User Flow (How It Works)", 9)

    flow_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.7))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = CARD_BG
    flow_box.line.color.rgb = CARD_BORDER
    tf = flow_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "5-STEP USER FLOW"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY
    p.space_after = Pt(14)

    steps = [
        "Step 1: Welcome & Login → User logs in via Google or Email.",
        "Step 2: Room Entry → User creates room or types room code.",
        "Step 3: Room Loading → App finds room code and opens room.",
        "Step 4: Live Sync Loop → Sofa shows members and live chat runs.",
        "Step 5: Leaving Room → User closes tab or clicks Leave Room."
    ]
    for s in steps:
        p = tf.add_paragraph()
        p.text = s
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

    uf_highlights = [
        ("Easy Login Check", "User logs in first, so their correct name and email show in member list."),
        ("Instant Code Lookup", "Room codes or invite links are checked quickly to join the exact room."),
        ("Auto Cleanup", "When a user leaves or closes tab, their avatar disappears automatically.")
    ]
    top_pos = 1.9
    for title, desc in uf_highlights:
        card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(top_pos), Inches(5.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)  # 12pt!
        p_desc.font.color.rgb = TEXT_MUTED

        top_pos += 1.65

    # ==================== SLIDE 10: Database Design ====================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Data Model", "Planned Database Structure", 10)

    schema_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.7))
    schema_box.fill.solid()
    schema_box.fill.fore_color.rgb = CARD_BG
    schema_box.line.color.rgb = CARD_BORDER
    tf = schema_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "ROOM DATABASE STRUCTURE"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY
    p.space_after = Pt(10)

    code_text = (
        "// Document: rooms/{roomId}\n"
        "{\n"
        "  \"roomId\": \"ayushmakvana-room\",\n"
        "  \"hostName\": \"Ayush\",\n"
        "  \"members\": [\n"
        "    { \"name\": \"Ayush\", \"isHost\": true },\n"
        "    { \"name\": \"Rahul\", \"isHost\": false }\n"
        "  ],\n"
        "  \"messages\": [\n"
        "    { \"sender\": \"Ayush\", \"text\": \"Hello!\" }\n"
        "  ]\n"
        "}"
    )
    p = tf.add_paragraph()
    p.text = code_text
    p.font.name = "Consolas"
    p.font.size = Pt(11)  # 11pt readable code
    p.font.color.rgb = TEXT_DARK

    db_highlights = [
        ("Room Code Table", "codes/{code} maps unique codes (AYUSHMAKVANA-2779) to the actual room."),
        ("NoSQL Realtime Storage", "Stores active sofa members and chat messages instantly in Google Cloud Firestore."),
        ("Memory Cache", "Keeps local data fast while preventing database lock crashes during page refresh.")
    ]
    top_pos = 1.9
    for title, desc in db_highlights:
        card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(top_pos), Inches(5.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)  # 12pt!
        p_desc.font.color.rgb = TEXT_MUTED

        top_pos += 1.65

    # ==================== SLIDE 11: Technology Stack ====================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Foundation", "Technology Stack", 11)

    t_w = Inches(3.7)
    t_h = Inches(2.2)

    techs = [
        ("FRONTEND", "React & Next.js 15 — modern website frame\nTypeScript — safer code & zero bug contracts"),
        ("BACKEND", "Node.js — fast server execution engine\nExpress.js — lightweight API route handler"),
        ("DATABASE & AUTH", "Firebase Auth — Google OAuth & Email login\nCloud Firestore — NoSQL real-time cloud DB"),
        ("DEVELOPMENT TOOLS", "Git & GitHub — code version history & team work\nTurbopack — ultra fast web development bundler"),
        ("NEXT MEDIA TOOLS", "YouTube Data API v3 — video search & details\nWebRTC & PeerJS — live camera & mic streaming"),
        ("CLOUD HOSTING", "Vercel / Firebase — fast & reliable web hosting")
    ]

    for idx, (cat, desc) in enumerate(techs):
        r = idx // 3
        c = idx % 3
        x = Inches(0.8 + c * 3.95)
        y = Inches(1.9 + r * 2.4)

        card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, t_w, t_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = cat
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = GREEN_PRIMARY
        p.space_after = Pt(6)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(12)  # 12pt!
        p_desc.font.color.rgb = TEXT_DARK

    # ==================== SLIDE 12: Development Roadmap ====================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Delivery", "Development Roadmap", 12)

    r_w = Inches(2.7)
    r_h = Inches(4.7)

    phases = [
        ("PHASE 1 (Completed)", "Planning & Architecture", "Requirement analysis, system design, tech selection, and database modeling.", "Time: 2 Weeks"),
        ("PHASE 2 (Completed)", "UI & Auth Creation", "Liquid glass UI design, Google 1-click login, Email signup, and profile menu.", "Time: 3 Weeks"),
        ("PHASE 3 (Completed)", "Room & Sync Engine", "Room code joining, real-time lounge sofa presence, and instant live chat.", "Time: 3 Weeks"),
        ("PHASE 4 (Next Phase)", "Video & WebRTC Integration", "YouTube API v3, synchronized video player, and camera/mic live stream.", "Time: 4 Weeks")
    ]

    for idx, (tag, title, desc, effort) in enumerate(phases):
        x = Inches(0.8 + idx * 2.95)
        card = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), r_w, r_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = tag
        p.font.name = FONT_MAIN
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = GREEN_PRIMARY
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(16)

        p = tf.add_paragraph()
        p.text = effort
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.bold = True
        p.font.color.rgb = GREEN_PRIMARY

    # ==================== SLIDE 13: Review 1 Progress ====================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "Delivery", "Review 1 Progress & Next Steps", 13)

    prog_w = Inches(5.6)
    prog_h = Inches(4.7)

    # Prepared Card
    c_prep = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), prog_w, prog_h)
    c_prep.fill.solid()
    c_prep.fill.fore_color.rgb = CARD_BG
    c_prep.line.color.rgb = CARD_BORDER
    tf = c_prep.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)

    p = tf.paragraphs[0]
    p.text = "COMPLETED IN REVIEW 1\nCurrent Working Features"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY
    p.space_after = Pt(12)

    preps = [
        "Frontend UI Design: Fully built Next.js 15 liquid-glass dark theme.",
        "Firebase Login: Google 1-click popup & Email login/signup connected.",
        "Room Join System: Room creation, code lookup (AYUSHMAKVANA-2779) & invite links working.",
        "Real-Time Sofa: Active room members show live on sofa with host crown icons.",
        "Live Chat Box: Fast instant text messaging working across all browsers."
    ]
    for item in preps:
        p = tf.add_paragraph()
        p.text = f"✓ {item}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Next Phase Card
    c_next = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), prog_w, prog_h)
    c_next.fill.solid()
    c_next.fill.fore_color.rgb = CARD_BG
    c_next.line.color.rgb = CARD_BORDER
    tf = c_next.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)

    p = tf.paragraphs[0]
    p.text = "UPCOMING IMPLEMENTATION\nNext Development Phase"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(12)

    nexts = [
        "YouTube API v3 Integration: Video search, playlist loading & details.",
        "Working Camera & Microphone: WebRTC live video & audio call stream.",
        "Global Video Sync: Play, pause & timestamp seek sync for all members in room.",
        "Host Controls: Host can enforce video playback sync so all stay together."
    ]
    for item in nexts:
        p = tf.add_paragraph()
        p.text = f"→ {item}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)  # 12pt!
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(10)

    # ==================== SLIDE 14: Thank You ====================
    slide14 = prs.slides.add_slide(blank_layout)
    set_bg(slide14)

    tb14 = slide14.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.733), Inches(2.0))
    p = tb14.text_frame.paragraphs[0]
    p.text = "Thank You !!"
    p.font.name = FONT_MAIN
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.alignment = PP_ALIGN.CENTER

    p_sub = tb14.text_frame.add_paragraph()
    p_sub.text = "Questions & Discussion"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(24)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.alignment = PP_ALIGN.CENTER

    ftb14 = slide14.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
    p = ftb14.text_frame.paragraphs[0]
    p.text = "SYNCSPACE · CHARUSAT UNIVERSITY"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GREEN_PRIMARY

    # Save PPTX to public directory
    output_path = r"c:\Users\HP\Downloads\Nextjs\syncspace\public\SyncSpace_Review_1_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
